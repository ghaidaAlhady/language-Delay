from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.rag.builder import build_knowledge_base
from app.rag.exceptions import KnowledgeBaseLoadError
from app.rag.loader import load_all_workbooks, load_workbook_sheets


def test_load_all_workbooks_succeeds_against_real_kb(kb_dir: Path) -> None:
    raw = load_all_workbooks(kb_dir)
    assert set(raw) == {"KB01.xlsx", "KB02.xlsx", "KB03.xlsx", "KB04.xlsx", "KB05.xlsx"}
    assert len(raw["KB01.xlsx"]["Age 2"]) == 75
    assert len(raw["KB02.xlsx"]["Activities"]) == 100
    assert len(raw["KB03.xlsx"]["Decision_Rules"]) == 64
    assert len(raw["KB05.xlsx"]["Assessment_Questions"]) == 80


def test_build_knowledge_base_succeeds_against_real_kb(kb_dir: Path) -> None:
    kb = build_knowledge_base(kb_dir)
    assert len(kb.milestones) == 300
    assert len(kb.references) == 3
    assert len(kb.activities) == 100
    assert len(kb.decision_rules) == 64
    assert len(kb.questions) == 80


def test_missing_kb_directory_raises() -> None:
    with pytest.raises(KnowledgeBaseLoadError):
        load_all_workbooks(Path("/no/such/directory"))


def test_missing_workbook_file_raises(tmp_path: Path) -> None:
    with pytest.raises(KnowledgeBaseLoadError, match="not found"):
        load_workbook_sheets(tmp_path, "KB01.xlsx")


def test_missing_required_sheet_raises(tmp_path: Path) -> None:
    # KB03 requires a "Decision_Rules" sheet; give it "Activities" instead.
    workbook = Workbook()
    workbook.active.title = "Activities"
    path = tmp_path / "KB03.xlsx"
    workbook.save(path)

    with pytest.raises(KnowledgeBaseLoadError, match="missing required sheet"):
        load_workbook_sheets(tmp_path, "KB03.xlsx")


def test_missing_required_column_raises(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Decision_Rules"
    sheet.append(["معرف القاعدة", "العمر"])  # missing most required columns
    sheet.append(["R001", 2])
    path = tmp_path / "KB03.xlsx"
    workbook.save(path)

    with pytest.raises(KnowledgeBaseLoadError, match="missing required column"):
        load_workbook_sheets(tmp_path, "KB03.xlsx")
